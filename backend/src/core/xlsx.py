"""Render plain rows into one .xlsx workbook (the project data export, #92).

Pure: no I/O beyond the in-memory buffer. Three things every cell gets:

- **Strings are always text cells.** openpyxl turns any string that starts with
  "=" into a formula, and these cells hold learner free text (assessment open
  answers, interview transcripts), so a typed "=HYPERLINK(...)" must never run
  when an admin opens the file.
- Control characters Excel's XML forbids are dropped (openpyxl raises on them),
  and text longer than Excel's 32,767-character cell limit is cut with a marker.
- Datetimes are written as naive UTC; Excel has no time zones.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from collections.abc import Sequence

__all__ = ["CELL_LIMIT", "Cell", "Sheet", "render_workbook"]

CELL_LIMIT = 32_767
_TRUNCATED = " … [truncated]"
_MIN_WIDTH, _MAX_WIDTH = 10, 60

type Cell = str | int | bool | datetime | None


@dataclass
class Sheet:
    title: str
    headers: Sequence[str]
    rows: list[Sequence[Cell]] = field(default_factory=list)


def _clean(value: Cell) -> Cell:
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone(UTC).replace(tzinfo=None)
        return value
    if isinstance(value, str):
        text = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(text) > CELL_LIMIT:
            text = text[: CELL_LIMIT - len(_TRUNCATED)] + _TRUNCATED
        return text
    return value


def render_workbook(sheets: Sequence[Sheet]) -> bytes:
    """One worksheet per Sheet, in order: bold frozen header row, then the rows."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        worksheet = workbook.create_sheet(title=sheet.title[:31])
        worksheet.append(list(sheet.headers))
        for header_cell in worksheet[1]:
            header_cell.font = Font(bold=True)
        widths = [len(header) for header in sheet.headers]
        for row in sheet.rows:
            worksheet.append([_clean(value) for value in row])
            for index, value in enumerate(row[: len(widths)]):
                if isinstance(value, str):
                    widths[index] = max(widths[index], min(len(value), _MAX_WIDTH))
        for excel_row in worksheet.iter_rows(min_row=2):
            for cell in excel_row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        worksheet.freeze_panes = "A2"
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = max(
                _MIN_WIDTH, min(width + 2, _MAX_WIDTH)
            )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
