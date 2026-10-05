"""Project data export — GET /api/admin/projects/{project_id}/export (#92).

End to end against real Postgres: the four batched repository reads, the scope
guard, the .xlsx response and the log line. Which values land in which column is
pinned by tests/unit/test_project_export_service.py on fake repositories."""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from typing import TYPE_CHECKING, Any

import pytest
import pytest_asyncio
from openpyxl import load_workbook

from src.repositories.conversation_repository import ConversationRepository
from src.repositories.message_repository import MessageRepository
from tests.fixtures.auth import auth_headers
from tests.fixtures.factories import (
    assign_lesson,
    complete_lesson,
    create_company,
    create_lesson,
    create_module,
    create_project,
    create_survey,
    create_track,
    create_user,
    make_project_admin,
)

if TYPE_CHECKING:
    from httpx import AsyncClient, Response
    from sqlalchemy.ext.asyncio import AsyncSession

    from src.core.configs.settings import Settings
    from src.db.models import Project, User

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class ExportSeed:
    project: Project
    other_project: Project
    power: User
    coach: User  # project admin of `project`, also a member of it -> staff, not a learner
    outsider_admin: User  # project admin of `other_project` only
    ann: User
    bob: User


@pytest_asyncio.fixture
async def seed(db_session: AsyncSession) -> ExportSeed:
    company = await create_company(db_session, "Export Co")
    project = await create_project(db_session, company, "Export Project")
    other_project = await create_project(db_session, company, "Export Other")
    power = await create_user(db_session, email="export-power@example.com", role="admin")

    def member(email: str, first: str, last: str, **kwargs: Any) -> Any:
        return create_user(
            db_session,
            email=email,
            first_name=first,
            last_name=last,
            company_id=company.id,
            project_id=project.id,
            **kwargs,
        )

    ann = await member("export-ann@example.com", "Ann", "Archer", password="hunter2hunter2")
    bob = await member("export-bob@example.com", "Bob", "Baker")
    coach = await member("export-coach@example.com", "Cora", "Coach", password="hunter2hunter2")
    await make_project_admin(db_session, coach, project, granted_by=power)
    outsider_admin = await create_user(
        db_session,
        email="export-outsider@example.com",
        password="hunter2hunter2",
        company_id=company.id,
        project_id=other_project.id,
    )
    await make_project_admin(db_session, outsider_admin, other_project, granted_by=power)
    await create_user(db_session, email="export-elsewhere@example.com", project_id=other_project.id)

    track = await create_track(db_session, "Export Track", required_level="practitioner")
    module = await create_module(db_session, track, "Export Module")
    first = await create_lesson(db_session, module, "Export L1", sort_order=0)
    second = await create_lesson(db_session, module, "Export L2", sort_order=1)
    await assign_lesson(db_session, ann, first)
    await assign_lesson(db_session, ann, second)
    await complete_lesson(db_session, ann, first, answers={"q1": "never exported"})

    await create_survey(
        db_session, ann, status="complete", score_level="practitioner", q_excitement="Less email"
    )
    await create_survey(db_session, bob)

    conversations = ConversationRepository(db_session)
    interview = await conversations.create(user_id=ann.id, kind="interview")
    await conversations.update_summary(
        interview.id,
        {
            "shared": {"where_you_are_today": "Early days"},
            "internal": {"company_context": "Retail"},
        },
    )
    await conversations.create(user_id=ann.id, kind="chatbot")
    messages = MessageRepository(db_session)
    await messages.create(interview.id, role="assistant", content="Welcome", metadata=None)
    await messages.create(interview.id, role="user", content="Thanks", metadata=None)
    return ExportSeed(project, other_project, power, coach, outsider_admin, ann, bob)


async def _export(
    client: AsyncClient, settings: Settings, caller: User, project: Project, query: str
) -> Response:
    return await client.get(
        f"/api/admin/projects/{project.id}/export?{query}",
        headers=auth_headers(caller, settings),
    )


def _rows(content: bytes) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(io.BytesIO(content))
    result = {}
    for name in workbook.sheetnames:
        rows = list(workbook[name].iter_rows(values_only=True))
        result[name] = [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]
    return result


_ALL = "sections=people&sections=assessment&sections=progress&sections=interview"


async def test_power_admin_downloads_every_section(
    client: AsyncClient, test_settings: Settings, seed: ExportSeed
) -> None:
    response = await _export(client, test_settings, seed.power, seed.project, _ALL)

    assert response.status_code == 200
    assert response.headers["content-type"] == _XLSX
    assert response.headers["cache-control"] == "no-store"
    assert 'filename="export-project-export-' in response.headers["content-disposition"]
    sheets = _rows(response.content)
    assert list(sheets) == [
        "People",
        "Assessment",
        "Progress",
        "Interview playbook",
        "Interview transcript",
    ]
    # Learners only: the coach (project admin), the power admin and the other
    # project's members are not in the file.
    assert [r["Email"] for r in sheets["People"]] == [seed.ann.email, seed.bob.email]
    assert [(r["Email"], r["Status"]) for r in sheets["Assessment"]] == [
        (seed.ann.email, "Complete"),
        (seed.bob.email, "Draft"),
    ]
    assert sheets["Assessment"][0]["What excites you"] == "Less email"
    assert [(r["Lesson"], r["Completed"]) for r in sheets["Progress"]] == [
        ("Export L1", "Yes"),
        ("Export L2", "No"),
    ]
    assert [(r["Part"], r["Value"]) for r in sheets["Interview playbook"]] == [
        ("shared", "Early days"),
        ("internal", "Retail"),
    ]
    assert [(r["Speaker"], r["Text"]) for r in sheets["Interview transcript"]] == [
        ("Interviewer", "Welcome"),
        ("Learner", "Thanks"),
    ]


async def test_project_admin_exports_their_own_project_only(
    client: AsyncClient, test_settings: Settings, seed: ExportSeed
) -> None:
    own = await _export(client, test_settings, seed.coach, seed.project, "sections=people")
    assert own.status_code == 200
    assert list(_rows(own.content)) == ["People"]

    foreign = await _export(
        client, test_settings, seed.outsider_admin, seed.project, "sections=people"
    )
    assert foreign.status_code == 404


@pytest.mark.parametrize("query", ["", "sections=people&sections=quiz", "sections=feedback"])
async def test_missing_or_unknown_sections_are_422(
    client: AsyncClient, test_settings: Settings, seed: ExportSeed, query: str
) -> None:
    response = await _export(client, test_settings, seed.power, seed.project, query)

    assert response.status_code == 422


async def test_the_log_line_names_ids_never_emails(
    client: AsyncClient,
    test_settings: Settings,
    seed: ExportSeed,
    caplog: pytest.LogCaptureFixture,
) -> None:
    with caplog.at_level(logging.INFO, logger="src.routers.admin_export"):
        response = await _export(client, test_settings, seed.coach, seed.project, _ALL)

    assert response.status_code == 200
    lines = [r.getMessage() for r in caplog.records if r.name == "src.routers.admin_export"]
    assert lines == [
        f"project export: project={seed.project.id} "
        f"sections=people,assessment,progress,interview by={seed.coach.id}"
    ]
