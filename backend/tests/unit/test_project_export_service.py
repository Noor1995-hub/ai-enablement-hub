"""ProjectExportService (#92) on fake repositories: who is in the file, which
sheets come out, and what their columns hold. The router, the repositories and
the authz row are covered by tests/integration/test_routers/test_admin_export.py."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

from src.repositories.base import (
    AbstractAssignmentRepository,
    AbstractAuthSessionRepository,
    AbstractContentRepository,
    AbstractConversationRepository,
    AbstractMessageRepository,
    AbstractOrgRepository,
    AbstractProgressRepository,
    AbstractProjectAdminRepository,
    AbstractSurveyRepository,
    AbstractUserRepository,
)
from src.services.exceptions import NotFoundError
from src.services.project_export_service import ProjectExportService

_T0 = datetime(2026, 9, 1, 9, 0, tzinfo=UTC)
_PROJECT = SimpleNamespace(id=uuid.uuid4(), name="Rollout Été 2026")


def _user(first: str, last: str, **overrides: Any) -> SimpleNamespace:
    values = {
        "id": uuid.uuid4(),
        "first_name": first,
        "last_name": last,
        "email": f"{first.lower()}@example.com",
        "role": "user",
        "is_active": True,
        "password_set": True,
        "external_id": None,
        "maturity_override": None,
        "created_at": _T0,
        "project_id": _PROJECT.id,
    }
    return SimpleNamespace(**(values | overrides))


def _survey(user: SimpleNamespace, **overrides: Any) -> SimpleNamespace:
    values: dict[str, Any] = {
        "user_id": user.id,
        "status": "complete",
        "started_at": _T0,
        "completed_at": _T0,
        "job_title": "Analyst",
        "department": "Finance",
        "score_level": "practitioner",
        "usage_band": "regular",
        "safety_flag": False,
        "b_freq": "WEEKLY",
        "b_breadth": "FEW",
        "b_tools": ["CLAUDE", "COPILOT"],
        "b_setup": None,
        "b_barriers": ["NO_TIME"],
        "b_attitude": 4,
        "b_readiness": 2,
        "q_excitement": "=SUM(A1:A9)",
        "q_timesink": "Reports",
        "c5_codes": ["KEPT_OUT"],
    }
    for n in range(1, 7):
        values[f"c{n}_rung"] = n % 4
        if n != 5:
            values[f"c{n}_code"] = "HAVE"
    return SimpleNamespace(**(values | overrides))


class _World:
    """The fake repositories, seeded with: ann (complete assessment, one done
    and one open lesson, a finished interview), bob (draft assessment, an
    unassigned Introduction lesson completed, an interview in progress),
    cid (invited, nothing), a power admin and a project admin (both excluded)."""

    def __init__(self) -> None:
        self.ann = _user("Ann", "Archer")
        self.bob = _user("Bob", "Baker", maturity_override="integrator")
        self.cid = _user("Cid", "Cole", password_set=False)
        self.power = _user("Pat", "Power", role="admin")
        self.staff = _user("Sam", "Staff")

        intro = SimpleNamespace(
            id=uuid.uuid4(), name="Introduction", sort_order=0, required_level="explore"
        )
        track = SimpleNamespace(
            id=uuid.uuid4(), name="General I", sort_order=1, required_level="practitioner"
        )
        self.intro_track, self.track = intro, track
        intro_module = SimpleNamespace(
            id=uuid.uuid4(), track_id=intro.id, title="Welcome", sort_order=0
        )
        module = SimpleNamespace(id=uuid.uuid4(), track_id=track.id, title="Basics", sort_order=0)
        self.l_intro = SimpleNamespace(
            id=uuid.uuid4(), module_id=intro_module.id, title="Hello", sort_order=0
        )
        self.l1 = SimpleNamespace(
            id=uuid.uuid4(), module_id=module.id, title="Prompts", sort_order=0
        )
        self.l2 = SimpleNamespace(
            id=uuid.uuid4(), module_id=module.id, title="Checks", sort_order=1
        )

        self.users = AsyncMock(spec=AbstractUserRepository)
        self.users.list_with_org_names.return_value = [
            (u, "Acme", _PROJECT.name)
            for u in (self.cid, self.power, self.bob, self.staff, self.ann)
        ]
        self.org = AsyncMock(spec=AbstractOrgRepository)
        self.org.get_project.return_value = _PROJECT
        self.project_admins = AsyncMock(spec=AbstractProjectAdminRepository)
        self.project_admins.admin_ids_for_project.return_value = {self.staff.id}
        self.sessions = AsyncMock(spec=AbstractAuthSessionRepository)
        self.sessions.last_used_by_user.return_value = {self.ann.id: _T0}
        self.surveys = AsyncMock(spec=AbstractSurveyRepository)
        self.surveys.list_for_users.return_value = [
            _survey(self.ann),
            _survey(self.bob, status="in_progress", completed_at=None, score_level=None),
        ]
        self.assignments = AsyncMock(spec=AbstractAssignmentRepository)
        self.assignments.assigned_lesson_counts_by_module_for_users.return_value = {
            self.ann.id: {module.id: (2, 1)}
        }
        self.assignments.user_lessons_for_users.return_value = [
            SimpleNamespace(user_id=self.ann.id, lesson_id=self.l2.id),
            SimpleNamespace(user_id=self.ann.id, lesson_id=self.l1.id),
        ]
        self.content = AsyncMock(spec=AbstractContentRepository)
        self.content.list_tracks.return_value = [intro, track]
        self.content.list_modules.return_value = [intro_module, module]
        self.content.list_lesson_summaries.return_value = [self.l_intro, self.l1, self.l2]
        self.progress = AsyncMock(spec=AbstractProgressRepository)
        self.progress.completions_for_users.return_value = {
            (self.ann.id, self.l1.id): _T0,
            (self.bob.id, self.l_intro.id): _T0,
        }
        done = SimpleNamespace(
            id=uuid.uuid4(),
            user_id=self.ann.id,
            language="nl",
            summary={
                "shared": {
                    "where_you_are_today": "Early",
                    "use_cases": [{"rank": 1, "name": "Mail"}],
                },
                "internal": {"blockers": {"stated": ["time"]}},
                "meta": {"model": "never exported"},
            },
        )
        open_ = SimpleNamespace(id=uuid.uuid4(), user_id=self.bob.id, language=None, summary=None)
        self.conversations = AsyncMock(spec=AbstractConversationRepository)
        self.conversations.interviews_for_users.return_value = [done, open_]
        self.messages = AsyncMock(spec=AbstractMessageRepository)
        self.messages.list_for_conversations.return_value = [
            SimpleNamespace(
                conversation_id=done.id,
                role="assistant",
                content="Hi! [[READY_TO_BUILD]]",
                created_at=_T0,
            ),
            SimpleNamespace(conversation_id=done.id, role="user", content="Hello", created_at=_T0),
            SimpleNamespace(
                conversation_id=open_.id, role="assistant", content="Hoi", created_at=_T0
            ),
        ]

    def service(self) -> ProjectExportService:
        return ProjectExportService(
            self.users,
            self.org,
            self.project_admins,
            self.sessions,
            self.surveys,
            self.assignments,
            self.content,
            self.progress,
            self.conversations,
            self.messages,
        )


async def _sheets(world: _World, sections: list[str]) -> dict[str, list[dict[str, Any]]]:
    export = await world.service().export(_PROJECT.id, sections)
    workbook = load_workbook(io.BytesIO(export.content))
    result = {}
    for name in workbook.sheetnames:
        rows = list(workbook[name].iter_rows(values_only=True))
        result[name] = [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]
    return result


async def test_sheets_follow_the_fixed_order_and_the_filename_is_a_dated_slug() -> None:
    world = _World()
    export = await world.service().export(_PROJECT.id, ["interview", "people"])

    workbook = load_workbook(io.BytesIO(export.content))
    assert workbook.sheetnames == ["People", "Interview playbook", "Interview transcript"]
    assert export.filename.startswith("rollout-ete-2026-export-")
    assert export.filename.endswith(".xlsx")


async def test_unknown_project_is_not_found() -> None:
    world = _World()
    world.org.get_project.return_value = None

    with pytest.raises(NotFoundError):
        await world.service().export(_PROJECT.id, ["people"])


async def test_people_lists_learners_only_sorted_with_status_and_maturity() -> None:
    world = _World()
    people = (await _sheets(world, ["people"]))["People"]

    assert [row["Email"] for row in people] == [
        "ann@example.com",
        "bob@example.com",
        "cid@example.com",
    ]
    ann, bob, cid = people
    assert (ann["Status"], ann["Maturity"], ann["Assessment"]) == (
        "Active",
        "practitioner",
        "Complete",
    )
    assert ann["Last active (UTC)"] == _T0.replace(tzinfo=None)
    # A draft never scores; the admin pin alone sets the base.
    assert (bob["Maturity"], bob["Maturity pinned by admin"], bob["Assessment"]) == (
        "integrator",
        "integrator",
        "Draft",
    )
    assert (cid["Status"], cid["Maturity"], cid["Assessment"]) == ("Invited", None, "Not started")


async def test_assessment_shows_labels_next_to_codes_and_flags_drafts() -> None:
    world = _World()
    rows = (await _sheets(world, ["assessment"]))["Assessment"]

    assert [row["Email"] for row in rows] == ["ann@example.com", "bob@example.com"]
    ann, bob = rows
    assert ann["Status"] == "Complete"
    assert bob["Status"] == "Draft"
    assert ann["C1 code"] == "HAVE"
    assert ann["C1 Understanding"].startswith("It predicts likely wording")
    assert ann["C5 Responsible use"] == "Mostly kept work information out of public tools"
    assert ann["C5 code"] == "KEPT_OUT"
    assert (ann["AI use frequency"], ann["AI use frequency code"]) == ("Most weeks", "WEEKLY")
    assert ann["Tools used"] == "Claude; Microsoft Copilot"
    assert ann["Tools used codes"] == "CLAUDE; COPILOT"
    assert ann["Things set up"] is None
    assert (ann["Attitude"], ann["Attitude code"]) == ("Positive", 4)
    assert ann["C6 rung"] == 2
    # Learner free text that looks like a formula is exported as text.
    assert ann["What excites you"] == "=SUM(A1:A9)"


async def test_progress_has_assigned_and_completed_lessons_in_outline_order() -> None:
    world = _World()
    rows = (await _sheets(world, ["progress"]))["Progress"]

    assert [(r["Email"], r["Lesson"], r["Assigned"], r["Completed"]) for r in rows] == [
        ("ann@example.com", "Prompts", "Yes", "Yes"),
        ("ann@example.com", "Checks", "Yes", "No"),
        ("bob@example.com", "Hello", "No", "Yes"),
    ]
    assert rows[0]["Track"] == "General I"
    assert rows[0]["Module"] == "Basics"
    assert rows[1]["Completed on (UTC)"] is None


async def test_interview_flattens_both_boxes_and_lists_every_turn() -> None:
    world = _World()
    sheets = await _sheets(world, ["interview"])

    playbook = sheets["Interview playbook"]
    assert [(r["Email"], r["Part"], r["Field"], r["Value"]) for r in playbook] == [
        ("ann@example.com", "shared", "where_you_are_today", "Early"),
        ("ann@example.com", "shared", "use_cases[1].rank", "1"),
        ("ann@example.com", "shared", "use_cases[1].name", "Mail"),
        ("ann@example.com", "internal", "blockers.stated[1]", "time"),
    ]
    assert {r["Language"] for r in playbook} == {"nl"}

    transcript = sheets["Interview transcript"]
    assert [
        (r["Email"], r["#"], r["Speaker"], r["Text"], r["Interview status"]) for r in transcript
    ] == [
        ("ann@example.com", 1, "Interviewer", "Hi!", "Complete"),
        ("ann@example.com", 2, "Learner", "Hello", "Complete"),
        ("bob@example.com", 1, "Interviewer", "Hoi", "In progress"),
    ]


async def test_an_unrequested_section_is_not_read() -> None:
    world = _World()
    await _sheets(world, ["people"])

    world.conversations.interviews_for_users.assert_not_awaited()
    world.progress.completions_for_users.assert_not_awaited()
