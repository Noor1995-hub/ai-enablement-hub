from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from src.core.xlsx import CELL_LIMIT, Sheet, render_workbook


def _load(content: bytes):  # noqa: ANN202 - openpyxl's Workbook is untyped
    return load_workbook(io.BytesIO(content))


def test_one_sheet_per_section_with_a_bold_frozen_header() -> None:
    workbook = _load(
        render_workbook([Sheet("People", ["Name", "Age"], [["Ann", 3]]), Sheet("Empty", ["A"])])
    )

    assert workbook.sheetnames == ["People", "Empty"]
    people = workbook["People"]
    assert [[cell.value for cell in row] for row in people.iter_rows()] == [
        ["Name", "Age"],
        ["Ann", 3],
    ]
    assert people["A1"].font.bold
    assert people.freeze_panes == "A2"
    assert workbook["Empty"].max_row == 1


def test_learner_text_that_looks_like_a_formula_stays_text() -> None:
    payload = '=HYPERLINK("https://evil.example","click")'
    sheet = _load(render_workbook([Sheet("S", ["Text"], [[payload]])]))["S"]

    assert sheet["A2"].value == payload
    assert sheet["A2"].data_type == "s"


def test_control_characters_are_dropped_and_huge_text_is_cut() -> None:
    sheet = _load(
        render_workbook([Sheet("S", ["Text"], [["a\x00b\x07c"], ["x" * (CELL_LIMIT + 10)]])])
    )["S"]

    assert sheet["A2"].value == "abc"
    assert len(sheet["A3"].value) == CELL_LIMIT
    assert sheet["A3"].value.endswith("[truncated]")


def test_aware_datetimes_are_written_as_naive_utc() -> None:
    brussels = datetime(2026, 10, 2, 14, 0, tzinfo=timezone(timedelta(hours=2)))
    sheet = _load(render_workbook([Sheet("S", ["At"], [[brussels], [None]])]))["S"]

    assert sheet["A2"].value == datetime(2026, 10, 2, 12, 0)
    assert sheet["A3"].value is None


def test_sheet_titles_are_cut_to_excels_31_characters() -> None:
    workbook = _load(render_workbook([Sheet("x" * 40, ["A"])]))

    assert workbook.sheetnames == ["x" * 31]
